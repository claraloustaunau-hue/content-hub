"""
Timeline PPT Generator

Takes an existing PPT timeline template and a list of events (JSON),
then generates a new PPT with all events in chronological order.

Usage:
    python timeline_generator.py --template template.pptx --events events.json --output output.pptx
    python timeline_generator.py --events events.json --output output.pptx
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

MAX_EVENTS_PER_SLIDE = 8

COLORS = {
    "line": RGBColor(0x33, 0x33, 0x33),
    "dot": RGBColor(0x00, 0x70, 0xC0),
    "date_text": RGBColor(0x00, 0x70, 0xC0),
    "event_text": RGBColor(0x33, 0x33, 0x33),
    "title_text": RGBColor(0xFF, 0xFF, 0xFF),
    "title_bg": RGBColor(0x00, 0x70, 0xC0),
    "slide_bg": RGBColor(0xFF, 0xFF, 0xFF),
}

FONT_NAME = "Calibri"


def parse_events(events_path: str) -> list[dict]:
    path = Path(events_path)
    if not path.exists():
        print(f"Error: events file '{events_path}' not found.")
        sys.exit(1)

    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, list):
        print("Error: events JSON must be a list of objects.")
        sys.exit(1)

    parsed = []
    for i, event in enumerate(data):
        if "date" not in event or "title" not in event:
            print(f"Error: event at index {i} must have 'date' and 'title' fields.")
            sys.exit(1)

        date_str = event["date"]
        try:
            date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        except ValueError:
            try:
                date_obj = datetime.strptime(date_str, "%Y-%m")
            except ValueError:
                try:
                    date_obj = datetime.strptime(date_str, "%Y")
                except ValueError:
                    print(
                        f"Error: event at index {i} has invalid date '{date_str}'. "
                        "Use YYYY-MM-DD, YYYY-MM, or YYYY."
                    )
                    sys.exit(1)

        parsed.append(
            {
                "date": date_obj,
                "date_str": event.get("display_date", date_str),
                "title": event["title"],
                "description": event.get("description", ""),
            }
        )

    parsed.sort(key=lambda e: e["date"])
    return parsed


def extract_events_from_pptx(pptx_path: str) -> list[dict]:
    """Try to extract text content from existing PPT slides for reference."""
    prs = Presentation(pptx_path)
    texts = []
    for slide in prs.slides:
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_texts.append(text)
        if slide_texts:
            texts.append(slide_texts)
    return texts


def add_title_bar(slide, title_text: str):
    left = Inches(0)
    top = Inches(0)
    width = SLIDE_WIDTH
    height = Inches(0.9)

    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["title_bg"]
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLORS["title_text"]
    p.font.name = FONT_NAME
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_timeline_line(slide, x_start, x_end, y_center):
    line = slide.shapes.add_connector(1, x_start, y_center, x_end, y_center)
    line.line.color.rgb = COLORS["line"]
    line.line.width = Pt(3)


def add_dot(slide, cx, cy, radius):
    left = cx - radius
    top = cy - radius
    diameter = radius * 2
    shape = slide.shapes.add_shape(9, left, top, diameter, diameter)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["dot"]
    shape.line.fill.background()


def add_text_box(slide, left, top, width, height, text, font_size, color, bold=False, alignment=PP_ALIGN.CENTER):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = font_size
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = FONT_NAME
    p.alignment = alignment
    return txBox


def build_timeline_slide(prs, events_chunk: list[dict], slide_title: str, slide_number: int, total_slides: int):
    slide_layout = prs.slide_layouts[6]  # blank layout
    slide = prs.slides.add_slide(slide_layout)

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLORS["slide_bg"]

    if total_slides > 1:
        title = f"{slide_title} ({slide_number}/{total_slides})"
    else:
        title = slide_title
    add_title_bar(slide, title)

    n = len(events_chunk)
    margin_x = Inches(1.0)
    timeline_y = Inches(4.2)
    x_start = margin_x
    x_end = SLIDE_WIDTH - margin_x
    timeline_width = x_end - x_start

    add_timeline_line(slide, x_start, x_end, timeline_y)

    arrow_size = Inches(0.15)
    arrow = slide.shapes.add_shape(
        7, x_end - Inches(0.05), timeline_y - arrow_size, arrow_size * 2, arrow_size * 2
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLORS["line"]
    arrow.line.fill.background()
    arrow.rotation = 90.0

    if n == 1:
        positions = [0.5]
    else:
        positions = [i / (n - 1) for i in range(n)]

    dot_radius = Inches(0.12)
    col_width = Inches(1.6)

    for idx, event in enumerate(events_chunk):
        frac = positions[idx]
        cx = x_start + int(timeline_width * frac)

        add_dot(slide, cx, timeline_y, dot_radius)

        vertical_line_len = Inches(0.5) if idx % 2 == 0 else Inches(0.5)
        is_above = idx % 2 == 0

        if is_above:
            vline_top = timeline_y - vertical_line_len - dot_radius
            vline_bottom = timeline_y - dot_radius
        else:
            vline_top = timeline_y + dot_radius
            vline_bottom = timeline_y + dot_radius + vertical_line_len

        connector = slide.shapes.add_connector(1, cx, vline_top, cx, vline_bottom)
        connector.line.color.rgb = COLORS["dot"]
        connector.line.width = Pt(1.5)

        date_box_height = Inches(0.35)
        event_box_height = Inches(1.2)

        if is_above:
            date_top = vline_top - date_box_height
            event_top = date_top - event_box_height
        else:
            date_top = vline_bottom + Inches(0.05)
            event_top = date_top + date_box_height + Inches(0.05)

        date_left = cx - col_width // 2
        add_text_box(
            slide, date_left, date_top, col_width, date_box_height,
            event["date_str"], Pt(11), COLORS["date_text"], bold=True
        )

        event_text = event["title"]
        if event["description"]:
            event_text += f"\n{event['description']}"

        add_text_box(
            slide, date_left, event_top, col_width, event_box_height,
            event_text, Pt(9), COLORS["event_text"], bold=False
        )


def generate_timeline(events: list[dict], output_path: str, template_path: str | None = None, title: str = "Línea de Tiempo"):
    if template_path:
        prs = Presentation(template_path)
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT
    else:
        prs = Presentation()
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT

    chunks = []
    for i in range(0, len(events), MAX_EVENTS_PER_SLIDE):
        chunks.append(events[i : i + MAX_EVENTS_PER_SLIDE])

    total_slides = len(chunks)
    for idx, chunk in enumerate(chunks):
        build_timeline_slide(prs, chunk, title, idx + 1, total_slides)

    prs.save(output_path)
    print(f"Timeline saved to '{output_path}' with {total_slides} slide(s) and {len(events)} event(s).")


def main():
    parser = argparse.ArgumentParser(
        description="Generate a timeline PPT from a list of events."
    )
    parser.add_argument(
        "--template",
        type=str,
        default=None,
        help="Path to an existing PPT to use as template (new slides will be appended).",
    )
    parser.add_argument(
        "--events",
        type=str,
        required=True,
        help="Path to a JSON file with events. Each event needs 'date' (YYYY-MM-DD) and 'title'.",
    )
    parser.add_argument(
        "--output",
        type=str,
        default="timeline_output.pptx",
        help="Output PPT file path (default: timeline_output.pptx).",
    )
    parser.add_argument(
        "--title",
        type=str,
        default="Línea de Tiempo",
        help="Title for the timeline slide(s).",
    )
    parser.add_argument(
        "--max-per-slide",
        type=int,
        default=8,
        help="Maximum number of events per slide (default: 8).",
    )

    args = parser.parse_args()

    global MAX_EVENTS_PER_SLIDE
    MAX_EVENTS_PER_SLIDE = args.max_per_slide

    events = parse_events(args.events)

    if not events:
        print("No events found in the JSON file.")
        sys.exit(1)

    print(f"Loaded {len(events)} events, sorted chronologically.")
    for e in events:
        print(f"  {e['date_str']}: {e['title']}")

    generate_timeline(events, args.output, args.template, args.title)


if __name__ == "__main__":
    main()
